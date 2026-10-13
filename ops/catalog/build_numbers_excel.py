"""Build Official WhatsApp numbers roster Excel from probe CSV."""

from __future__ import annotations

import csv
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
CSV_PATH = Path(__file__).with_name("_numbers_roster.raw")
OUT = ROOT / "Official-WhatsApp-Numbers-Roster.xlsx"

FONT = Font(name="Arial", size=11)
FONT_BOLD = Font(name="Arial", size=11, bold=True)
FONT_TITLE = Font(name="Arial", size=14, bold=True)
FONT_HEADER = Font(name="Arial", size=11, bold=True, color="FFFFFF")
FILL_HEADER = PatternFill("solid", fgColor="1F4E79")
FILL_TOTAL = PatternFill("solid", fgColor="D6EAF8")
FILL_UNASSIGNED = PatternFill("solid", fgColor="FDEBD0")
FILL_INACTIVE = PatternFill("solid", fgColor="F2F3F4")
THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)


def load_rows() -> list[dict]:
    text = CSV_PATH.read_text(encoding="utf-8", errors="replace")
    started = False
    lines = []
    for ln in text.splitlines():
        if ln.startswith("phone_number,"):
            started = True
        if not started:
            continue
        if ln.strip() == "PROBE_OK":
            break
        if ln.strip():
            lines.append(ln)
    reader = csv.DictReader(lines)
    return list(reader)


def style_header(ws, cols: int, row: int = 1):
    for col in range(1, cols + 1):
        cell = ws.cell(row, col)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN


def autosize(ws, min_w=10, max_w=40):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min_w
        for cell in col:
            if cell.value is None:
                continue
            width = max(width, min(max_w, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


def main():
    rows = load_rows()
    active = [r for r in rows if r["is_active"] == "yes"]
    inactive = [r for r in rows if r["is_active"] != "yes"]
    assigned = [r for r in active if r["agent_name"].strip()]
    unassigned = [r for r in active if not r["agent_name"].strip()]

    by_agent: dict[str, list] = defaultdict(list)
    for r in assigned:
        by_agent[r["agent_name"].strip()].append(r)

    wb = Workbook()

    # --- Summary ---
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Official WhatsApp Numbers Roster"
    ws["A1"].font = FONT_TITLE
    ws["A2"] = f"Source: live Official inbox (catalog_wabanumber) — {date.today().isoformat()}"
    ws["A2"].font = Font(name="Arial", size=10, italic=True)
    ws["A3"] = "One number can appear once; assignment is the active primary/shared roster on Numbers screen."
    ws["A3"].font = Font(name="Arial", size=10, italic=True)

    ws["A5"] = "Metric"
    ws["B5"] = "Count"
    style_header(ws, 2, 5)
    metrics = [
        ("Total numbers in database", len(rows)),
        ("Active numbers", len(active)),
        ("Inactive / hidden numbers", len(inactive)),
        ("Active - assigned to an agent", len(assigned)),
        ("Active - unassigned", len(unassigned)),
        ("Agents who have at least one number", len(by_agent)),
    ]
    for i, (label, value) in enumerate(metrics, start=6):
        ws.cell(i, 1, label).font = FONT
        ws.cell(i, 2, value).font = FONT_BOLD
        ws.cell(i, 1).border = THIN
        ws.cell(i, 2).border = THIN
        if "unassigned" in label.lower():
            ws.cell(i, 1).fill = FILL_UNASSIGNED
            ws.cell(i, 2).fill = FILL_UNASSIGNED

    ws["A13"] = "See sheet By Agent for who has which number; All Numbers for the full list."
    ws["A13"].font = Font(name="Arial", size=10, italic=True)

    autosize(ws)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 12

    # --- By Agent ---
    ws2 = wb.create_sheet("By Agent")
    headers = ["Agent", "Bitrix user id", "Numbers count", "Phone numbers", "Number labels", "Communities"]
    ws2.append(headers)
    style_header(ws2, len(headers))
    for agent in sorted(by_agent.keys(), key=str.lower):
        items = sorted(by_agent[agent], key=lambda x: x["phone_number"])
        phones = ", ".join(r["phone_number"] for r in items)
        labels = ", ".join((r["number_label"] or r["phone_number"]) for r in items)
        communities = ", ".join(sorted({r["community"] for r in items if r["community"]}))
        bitrix = items[0]["bitrix_user_id"]
        row_i = [
            agent,
            bitrix,
            len(items),
            phones,
            labels,
            communities,
        ]
        ws2.append(row_i)
        r = ws2.max_row
        for c in range(1, 7):
            cell = ws2.cell(r, c)
            cell.font = FONT
            cell.border = THIN
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Unassigned block
    ws2.append([])
    start = ws2.max_row + 1
    ws2.cell(start, 1, "UNASSIGNED (active)").font = FONT_BOLD
    ws2.cell(start, 1).fill = FILL_UNASSIGNED
    ws2.append(["(no agent)", "", len(unassigned), ", ".join(r["phone_number"] for r in unassigned), ", ".join((r["number_label"] or r["phone_number"]) for r in unassigned), ""])
    for c in range(1, 7):
        cell = ws2.cell(ws2.max_row, c)
        cell.font = FONT
        cell.border = THIN
        cell.fill = FILL_UNASSIGNED
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    autosize(ws2)
    ws2.column_dimensions["D"].width = 55
    ws2.column_dimensions["E"].width = 45
    ws2.row_dimensions[1].height = 22

    # --- All Numbers ---
    ws3 = wb.create_sheet("All Numbers")
    headers3 = [
        "Phone number",
        "Label",
        "Display name",
        "Community",
        "Account",
        "Active",
        "Webhook configured",
        "Source",
        "Assigned agent",
        "Agent username",
        "Bitrix user id",
        "Role",
        "Status",
    ]
    ws3.append(headers3)
    style_header(ws3, len(headers3))

    def status_of(r):
        if r["is_active"] != "yes":
            return "inactive"
        if r["agent_name"].strip():
            return "assigned"
        return "unassigned"

    ordered = sorted(
        rows,
        key=lambda r: (
            0 if r["is_active"] == "yes" else 1,
            0 if r["agent_name"].strip() else 1,
            (r["agent_name"] or "").lower(),
            r["phone_number"],
        ),
    )
    for r in ordered:
        st = status_of(r)
        ws3.append(
            [
                r["phone_number"],
                r["number_label"],
                r["display_name"],
                r["community"],
                r["account"],
                r["is_active"],
                r["webhook_configured"],
                r["source"],
                r["agent_name"],
                r["agent_username"],
                r["bitrix_user_id"],
                r["role"],
                st,
            ]
        )
        rr = ws3.max_row
        fill = None
        if st == "unassigned":
            fill = FILL_UNASSIGNED
        elif st == "inactive":
            fill = FILL_INACTIVE
        for c in range(1, 14):
            cell = ws3.cell(rr, c)
            cell.font = FONT
            cell.border = THIN
            if fill:
                cell.fill = fill

    last_data = ws3.max_row
    total_row = last_data + 2
    # Values (not formulas): LibreOffice recalc is unavailable on this Windows host.
    # Counts match the Summary sheet and the live export used to build this file.
    totals = [
        ("TOTAL ROWS", len(rows)),
        ("Active", len(active)),
        ("Assigned", len(assigned)),
        ("Unassigned (active)", len(unassigned)),
        ("Inactive", len(inactive)),
    ]
    col = 1
    for label, value in totals:
        ws3.cell(total_row, col, label).font = FONT_BOLD
        ws3.cell(total_row, col).fill = FILL_TOTAL
        ws3.cell(total_row, col).border = THIN
        ws3.cell(total_row, col + 1, value).font = FONT_BOLD
        ws3.cell(total_row, col + 1).fill = FILL_TOTAL
        ws3.cell(total_row, col + 1).border = THIN
        col += 2

    autosize(ws3)
    ws3.auto_filter.ref = f"A1:M{last_data}"
    ws3.freeze_panes = "A2"

    # --- Note ---
    ws4 = wb.create_sheet("Notes")
    ws4["A1"] = "Notes"
    ws4["A1"].font = FONT_TITLE
    notes = [
        "Live snapshot from Hostinger Official WhatsApp (catalog_wabanumber + active NumberAssignment).",
        "Orange rows = active but not assigned to any agent.",
        "Grey rows = inactive (hidden from normal assign / inbox use).",
        "Oct-8 Coastline / Blue Brick numbers (971581058717–726) are active unassigned until you assign them.",
        "Webhook column reflects our DB flag; paste the Official webhook in WBBOX if inbound is needed.",
        "Webhook URL: https://wa.srv1802669.hstgr.cloud/api/gateway/webhook/alots/",
    ]
    for i, line in enumerate(notes, start=3):
        ws4.cell(i, 1, line).font = FONT
    ws4.column_dimensions["A"].width = 110

    wb.save(OUT)
    print(f"wrote {OUT}")
    print(f"total={len(rows)} active={len(active)} assigned={len(assigned)} unassigned={len(unassigned)} agents={len(by_agent)}")


if __name__ == "__main__":
    main()
