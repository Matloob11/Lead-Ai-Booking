"""Extract phones from a7jobs applications xlsx → campaign upload file."""
from __future__ import annotations

import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SRC = r"C:\Users\Matloob\Desktop\a7-full-whatsapp-system\a7jobs-applications-expired (3).xlsx"
OUT_XLSX = r"C:\Users\Matloob\Desktop\a7-full-whatsapp-system\a7jobs-interview-campaign-phones.xlsx"
OUT_CSV = r"C:\Users\Matloob\Desktop\a7-full-whatsapp-system\a7jobs-interview-campaign-phones.csv"


def digits_only(value) -> str:
    return "".join(c for c in str(value or "") if c.isdigit())


def normalize_phone(raw, default_cc: str = "971") -> str | None:
    text = str(raw or "").strip().replace("\u200e", "").replace("\u200f", "")
    if not text:
        return None
    digits = digits_only(text)
    if not digits:
        return None
    if default_cc == "971" and digits.startswith("05") and len(digits) == 10:
        digits = "971" + digits[1:]
    elif default_cc == "971" and digits.startswith("5") and len(digits) == 9:
        digits = "971" + digits
    if len(digits) < 8 or len(digits) > 15:
        return None
    return digits


def main() -> None:
    wb = load_workbook(SRC, read_only=True, data_only=True)
    ws = wb["Applications"]
    rows = ws.iter_rows(values_only=True)
    header = [str(c or "").strip() for c in next(rows)]
    idx = {h: i for i, h in enumerate(header)}
    i_name = idx["Full Name"]
    i_m1 = idx["Mobile 1"]
    i_m2 = idx["Mobile 2"]
    i_job = idx["Job Listing"]

    seen: set[str] = set()
    recipients: list[tuple[str, str, str, str, str]] = []
    skipped = {"empty": 0, "invalid": 0, "dup": 0}
    total_raw = 0

    for row in rows:
        name = str(row[i_name] or "").strip()
        job = str(row[i_job] or "").strip()
        for label, ii in (("Mobile 1", i_m1), ("Mobile 2", i_m2)):
            total_raw += 1
            raw = row[ii] if ii < len(row) else None
            if raw is None or str(raw).strip() == "":
                skipped["empty"] += 1
                continue
            phone = normalize_phone(raw)
            if not phone:
                skipped["invalid"] += 1
                continue
            if phone in seen:
                skipped["dup"] += 1
                continue
            seen.add(phone)
            recipients.append((name or phone, phone, job, label, str(raw).strip()))

    wb.close()

    font = Font(name="Arial", size=11)
    head_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1C5F91")
    thin = Border(
        left=Side(style="thin", color="C8D2DC"),
        right=Side(style="thin", color="C8D2DC"),
        top=Side(style="thin", color="C8D2DC"),
        bottom=Side(style="thin", color="C8D2DC"),
    )

    ow = Workbook()
    # Active sheet must be Campaign — panel reads workbook.active for .xlsx
    camp = ow.active
    camp.title = "Campaign"
    camp["A1"] = "{{1}}"
    camp["B1"] = "Number"
    for cell in (camp["A1"], camp["B1"]):
        cell.font = head_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin

    for i, (name, phone, _job, _src, _raw) in enumerate(recipients, start=2):
        camp[f"A{i}"] = name
        camp[f"B{i}"] = "+" + phone
        camp[f"A{i}"].font = font
        camp[f"B{i}"].font = font
        camp[f"A{i}"].border = thin
        camp[f"B{i}"].border = thin

    camp.column_dimensions["A"].width = 28
    camp.column_dimensions["B"].width = 18

    info = ow.create_sheet("How to use")
    info["A1"] = "A7 WhatsApp campaign upload — interview applicants"
    info["A1"].font = Font(name="Arial", size=14, bold=True)
    lines = [
        "",
        "1. Open Campaigns in the WA panel.",
        "2. Upload a7jobs-interview-campaign-phones.xlsx (or the .csv).",
        "3. Select your interview WhatsApp template.",
        "4. Start the campaign.",
        "",
        "Required headers: {{1}} and Number",
        "{{1}} = applicant Full Name (template body {{1}} / Owner Data title).",
        "Number = +971… phone (E.164).",
        "",
        f"Unique phones: {len(recipients)}",
        "Source: a7jobs-applications-expired (3).xlsx → Applications (Mobile 1 + Mobile 2, deduped).",
        (
            f"Skipped empty={skipped['empty']} "
            f"invalid={skipped['invalid']} duplicate={skipped['dup']}"
        ),
        "",
        "If the template needs {{2}}, {{3}}, … add those columns in row 1 and fill every row.",
    ]
    for i, line in enumerate(lines, start=2):
        info[f"A{i}"] = line
        info[f"A{i}"].font = font
    info.column_dimensions["A"].width = 100

    det = ow.create_sheet("Source detail")
    det.append(["{{1}}", "Number", "Job Listing", "From column", "Raw phone"])
    for c in det[1]:
        c.font = head_font
        c.fill = fill
    for name, phone, job, src_col, raw in recipients:
        det.append([name, "+" + phone, job, src_col, raw])
    det.column_dimensions["A"].width = 28
    det.column_dimensions["B"].width = 18
    det.column_dimensions["C"].width = 55
    det.column_dimensions["D"].width = 12
    det.column_dimensions["E"].width = 18

    ow.save(OUT_XLSX)

    with open(OUT_CSV, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["{{1}}", "Number"])
        for name, phone, _job, _src, _raw in recipients:
            w.writerow([name, "+" + phone])

    uae = sum(1 for r in recipients if r[1].startswith("971"))
    print("unique", len(recipients))
    print("uae_971", uae, "other", len(recipients) - uae)
    print("skipped", skipped)
    print("total_raw_cells", total_raw)
    print("xlsx", OUT_XLSX)
    print("csv", OUT_CSV)
    print("sample", recipients[:3])


if __name__ == "__main__":
    main()
