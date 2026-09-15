"""Normalize and extract phone numbers for campaign recipients."""

from __future__ import annotations

import csv
import io
import re

# Digits runs that look like international mobiles (UAE 971…, etc.).
_DIGIT_RUN = re.compile(r"\d{8,15}")
_VAR_HEADER = re.compile(r"^\{\{\s*(\d+)\s*\}\}$", re.IGNORECASE)


class SpreadsheetFormatError(ValueError):
    """Fixed campaign spreadsheet layout was not followed."""


def digits_only(value: str) -> str:
    return "".join(ch for ch in str(value or "") if ch.isdigit())


def normalize_phone(raw: str, default_cc: str = "971") -> str | None:
    """Return E.164-ish digits without +, or None if unusable."""
    text = (raw or "").strip()
    if not text:
        return None
    # Strip common spreadsheet junk
    text = text.replace("\u200e", "").replace("\u200f", "")
    digits = digits_only(text)
    if not digits:
        return None
    # Local UAE mobile 05xxxxxxxx → 9715xxxxxxxx
    if default_cc == "971" and digits.startswith("05") and len(digits) == 10:
        digits = "971" + digits[1:]
    elif default_cc == "971" and digits.startswith("5") and len(digits) == 9:
        digits = "971" + digits
    if len(digits) < 8 or len(digits) > 15:
        return None
    return digits


def normalize_spreadsheet_header(raw) -> str:
    """Map a spreadsheet header cell to {{n}} or Number."""
    text = str(raw or "").strip()
    if not text:
        return ""
    match = _VAR_HEADER.match(text.replace(" ", ""))
    if match:
        return f"{{{{{match.group(1)}}}}}"
    if text.lower() == "number":
        return "Number"
    return text


def is_spreadsheet_workbook(file_bytes: bytes, filename: str = "") -> bool:
    """True when the upload is Excel (.xlsx/.xlsm), not plain CSV text."""
    name = (filename or "").lower()
    return name.endswith(".xlsx") or name.endswith(".xlsm") or (
        bool(file_bytes) and file_bytes[:2] == b"PK"
    )


def _read_spreadsheet_rows(file_bytes: bytes, filename: str = "") -> list[list[str]]:
    """Return all rows as strings from .xlsx or CSV."""
    name = (filename or "").lower()
    is_xlsx = is_spreadsheet_workbook(file_bytes, filename)
    if is_xlsx:
        try:
            import openpyxl
        except ImportError as exc:
            raise SpreadsheetFormatError(
                "Excel (.xlsx) support is not installed on the server."
            ) from exc
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        sheet = workbook.active
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(["" if cell is None else str(cell).strip() for cell in row])
        workbook.close()
        return rows

    text = file_bytes.decode("utf-8-sig", errors="replace")
    try:
        sample = text[:2048]
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    return [[str(cell or "").strip() for cell in row] for row in reader]


def _trim_trailing_blank_rows(rows: list[list[str]]) -> list[list[str]]:
    out = list(rows)
    while out and all(not str(cell or "").strip() for cell in out[-1]):
        out.pop()
    return out


def parse_campaign_spreadsheet(
    file_bytes: bytes,
    *,
    filename: str = "",
    body_var_count: int = 0,
) -> tuple[list[dict], list[str]]:
    """Parse the fixed A7 campaign Excel layout.

    Row 1 must name columns with ``{{1}}``, ``{{2}}``, … and ``Number``.
    Column order does not matter — only the header names. Each data row fills
    template body placeholders from the matching ``{{n}}`` column and the phone
    from ``Number``. ``{{1}}`` also becomes the Owner Data title when created.
    """
    rows = _trim_trailing_blank_rows(_read_spreadsheet_rows(file_bytes, filename))
    if not rows:
        raise SpreadsheetFormatError("Wrong format: the file is empty.")

    headers = [normalize_spreadsheet_header(cell) for cell in rows[0]]
    if not any(headers):
        raise SpreadsheetFormatError("Wrong format: row 1 must be the header row.")

    if "Number" not in headers:
        raise SpreadsheetFormatError(
            'Wrong format: row 1 must include a "Number" column for phone numbers.'
        )
    if "{{1}}" not in headers:
        raise SpreadsheetFormatError('Wrong format: row 1 must include a "{{1}}" column.')

    for number in range(1, body_var_count + 1):
        key = f"{{{{{number}}}}}"
        if key not in headers:
            raise SpreadsheetFormatError(
                f'Wrong format: this template needs a "{key}" column in row 1.'
            )

    number_idx = headers.index("Number")
    name_idx = headers.index("{{1}}")
    var_indices = {n: headers.index(f"{{{{{n}}}}}") for n in range(1, body_var_count + 1)}

    recipients: list[dict] = []
    warnings: list[str] = []
    seen: set[str] = set()

    for row_num, row in enumerate(rows[1:], start=2):
        if not row or all(not str(cell or "").strip() for cell in row):
            continue

        raw_phone = row[number_idx] if number_idx < len(row) else ""
        phone = normalize_phone(raw_phone)
        if not phone:
            warnings.append(f"Row {row_num}: invalid or missing phone in Number column.")
            continue
        if phone in seen:
            warnings.append(f"Row {row_num}: duplicate phone +{phone} — skipped.")
            continue

        variables: dict[str, str] = {}
        row_ok = True
        for number in range(1, body_var_count + 1):
            idx = var_indices[number]
            value = str(row[idx] if idx < len(row) else "").strip()
            if not value:
                warnings.append(f"Row {row_num}: {{{{{number}}}}} is empty.")
                row_ok = False
            else:
                variables[str(number)] = value

        if not row_ok:
            continue

        seen.add(phone)
        name = str(row[name_idx] if name_idx < len(row) else "").strip()
        recipients.append({"phone": phone, "name": name[:200], "variables": variables})

    if not recipients:
        detail = "Wrong format: no valid recipient rows found."
        if warnings:
            detail += " " + " ".join(warnings[:6])
        raise SpreadsheetFormatError(detail)

    return recipients, warnings


def parse_recipients_text(blob: str) -> list[dict]:
    """Parse pasted text: one phone per line, optional `phone,name` or `name,phone`."""
    out: list[dict] = []
    seen: set[str] = set()
    for line in (blob or "").replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        phone = None
        name = ""
        if "," in line or "\t" in line:
            parts = re.split(r"[,;\t]", line, maxsplit=1)
            a = (parts[0] or "").strip()
            b = (parts[1] or "").strip() if len(parts) > 1 else ""
            pa, pb = normalize_phone(a), normalize_phone(b)
            if pa and not pb:
                phone, name = pa, b if not digits_only(b) == pa else ""
            elif pb and not pa:
                phone, name = pb, a
            elif pa and pb:
                # Prefer the longer (more international) as phone; other as name if non-digit-heavy
                phone = pa if len(pa) >= len(pb) else pb
                other = b if phone == pa else a
                name = other if not normalize_phone(other) else ""
            else:
                # Fall through: scan line for a digit run
                pass
        if not phone:
            phone = normalize_phone(line)
            if not phone:
                m = _DIGIT_RUN.search(line)
                if m:
                    phone = normalize_phone(m.group(0))
        if not phone or phone in seen:
            continue
        seen.add(phone)
        out.append({"phone": phone, "name": (name or "")[:200]})
    return out


def parse_recipients_csv(file_bytes: bytes) -> list[dict]:
    """Parse CSV upload. Looks for phone/mobile/whatsapp and name/title columns."""
    text = file_bytes.decode("utf-8-sig", errors="replace")
    # Sniff dialect lightly
    try:
        sample = text[:2048]
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(io.StringIO(text), dialect)
    rows = list(reader)
    if not rows:
        return []

    header = [str(c or "").strip().lower() for c in rows[0]]
    phone_idx = None
    name_idx = None
    phone_keys = ("phone", "mobile", "whatsapp", "wa", "number", "msisdn", "contact")
    name_keys = ("name", "title", "owner", "contact name", "full name", "client")

    for i, h in enumerate(header):
        if phone_idx is None and any(k in h for k in phone_keys):
            phone_idx = i
        if name_idx is None and any(k == h or k in h for k in name_keys):
            name_idx = i

    start = 1
    if phone_idx is None:
        # No header — treat every cell as possible phone
        start = 0
        phone_idx = 0
        name_idx = 1 if len(header) > 1 else None

    out: list[dict] = []
    seen: set[str] = set()
    for row in rows[start:]:
        if not row:
            continue
        raw_phone = row[phone_idx] if phone_idx < len(row) else ""
        phone = normalize_phone(raw_phone)
        if not phone:
            # Scan row
            for cell in row:
                phone = normalize_phone(cell)
                if phone:
                    break
        if not phone or phone in seen:
            continue
        seen.add(phone)
        name = ""
        if name_idx is not None and name_idx < len(row):
            name = str(row[name_idx] or "").strip()
            if normalize_phone(name) == phone:
                name = ""
        out.append({"phone": phone, "name": name[:200]})
    return out
